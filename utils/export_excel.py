# utils/export_excel.py
import io
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def exportar_excel(dataframe: pd.DataFrame, titulo_extra: str = "",
                   edificio: str | None = None, laboratorio: str | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora D5"

    # Estilos
    azul_inst = "1F4E78"
    verde_ok   = "E8F5E9"
    rojo_falla = "FFEBEE"

    f_titulo   = Font(name="Calibri", size=14, bold=True, color=azul_inst)
    f_sub      = Font(name="Calibri", size=9,  italic=True, color="595959")
    f_header   = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    f_body     = Font(name="Calibri", size=10)
    f_bold     = Font(name="Calibri", size=10, bold=True)
    fill_hdr   = PatternFill(start_color=azul_inst, end_color=azul_inst, fill_type="solid")
    fill_zebra = PatternFill(start_color="F4F7FB", end_color="F4F7FB", fill_type="solid")
    fill_falla = PatternFill(start_color=rojo_falla, end_color=rojo_falla, fill_type="solid")
    th = Side(border_style="thin", color="BDC3C7")
    borde = Border(left=th, right=th, top=th, bottom=th)

    # ── Encabezado institucional dinámico ────────────────────────────
    if edificio:
        titulo_edif = f"DIVISIÓN DE TECNOLOGÍAS DE LA INFORMACIÓN — EDIFICIO {edificio}"
    else:
        titulo_edif = "DIVISIÓN DE TECNOLOGÍAS DE LA INFORMACIÓN — EDIFICIOS D5 Y K5"

    ws["B2"] = titulo_edif
    ws["B2"].font = f_titulo
    ws["B3"] = "Bitácora Digital de Uso de Laboratorios  |  Solo para fines educativos o de apoyo académico"
    ws["B3"].font = f_sub
    if titulo_extra:
        ws["B4"] = titulo_extra
        ws["B4"].font = f_sub

    ws["B6"] = "Laboratorio:"; ws["B6"].font = f_bold
    if laboratorio:
        labs_txt = f"Lab {laboratorio} — Edificio {edificio or ''}"
    elif edificio:
        from config.settings import EDIFICIOS
        labs_txt = f"Edificio {edificio}: Labs {', '.join(EDIFICIOS.get(edificio, []))}"
    else:
        from config.settings import LABORATORIOS
        labs_txt = f"Todos los laboratorios ({', '.join(LABORATORIOS)})"
    ws["C6"] = labs_txt; ws["C6"].font = f_body
    ws["E6"] = "Exportado:"; ws["E6"].font = f_bold
    ws["F6"] = datetime.now().strftime("%Y-%m-%d %H:%M"); ws["F6"].font = f_body

    headers = [
        "Fecha", "Turno", "Laboratorio", "Entrada", "Salida",
        "Grado/Grupo", "# Alumnos", "Asignatura", "Tema",
        "Nombre del Profesor", "Firma (Hash)", "Jefe de Grupo",
        "Observaciones", "Estado",
    ]

    fila_hdr = 8
    for ci, h in enumerate(headers, start=2):
        cell = ws.cell(row=fila_hdr, column=ci, value=h)
        cell.font = fill_hdr and f_header
        cell.fill = fill_hdr
        cell.border = borde
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila_hdr].height = 30

    fila_data = fila_hdr + 1
    col_estado_map = {col: i + 2 for i, col in enumerate(headers)}

    for idx, row in dataframe.iterrows():
        tiene_falla = str(row.get("tiene_falla", 0)) == "1"
        valores = [
            row.get("fecha", ""),         row.get("turno", ""),
            row.get("laboratorio", ""),   row.get("entrada", ""),
            row.get("salida", ""),        row.get("grado_grupo", ""),
            row.get("no_alum", ""),       row.get("asignatura", ""),
            row.get("tema", ""),          row.get("nombre_profesor", ""),
            row.get("firma_hash", ""),    row.get("firma_jefe_grupo", ""),
            row.get("observaciones", ""), row.get("analisis_obs", ""),
        ]
        for ci, val in enumerate(valores, start=2):
            cell = ws.cell(row=fila_data, column=ci, value=val)
            cell.font = f_body
            cell.border = borde
            cell.alignment = Alignment(
                horizontal="center" if ci <= 7 else "left",
                vertical="center", wrap_text=True
            )
            if tiene_falla:
                cell.fill = fill_falla
            elif idx % 2 == 0:
                cell.fill = fill_zebra
        ws.row_dimensions[fila_data].height = 20
        fila_data += 1

    # Firma de responsable
    fila_firma = fila_data + 2
    ws.cell(row=fila_firma,     column=2).value = "________________________________________________"
    ws.cell(row=fila_firma,     column=2).font  = f_body
    ws.cell(row=fila_firma + 1, column=2).value = "NOMBRE Y FIRMA DEL LABORATORISTA RESPONSABLE"
    ws.cell(row=fila_firma + 1, column=2).font  = f_bold

    # Ajustar anchos
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row >= fila_hdr and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        if col[0].column >= 2:
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 45), 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

def exportar_excel_inventario(dataframe: pd.DataFrame, titulo_extra: str = "") -> bytes:
    """
    Genera un Excel profesional del inventario con el mismo estilo
    institucional que exportar_excel() usa para la bitácora.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario D5"

    azul_inst  = "1F4E78"
    verde_ok   = "E8F5E9"
    amarillo   = "FFF8E1"
    rojo_malo  = "FFEBEE"

    f_titulo  = Font(name="Calibri", size=14, bold=True,  color=azul_inst)
    f_sub     = Font(name="Calibri", size=9,  italic=True, color="595959")
    f_header  = Font(name="Calibri", size=10, bold=True,  color="FFFFFF")
    f_body    = Font(name="Calibri", size=10)
    f_bold    = Font(name="Calibri", size=10, bold=True)
    fill_hdr  = PatternFill(start_color=azul_inst, end_color=azul_inst, fill_type="solid")
    fill_bueno   = PatternFill(start_color=verde_ok,  end_color=verde_ok,  fill_type="solid")
    fill_regular = PatternFill(start_color=amarillo,  end_color=amarillo,  fill_type="solid")
    fill_malo    = PatternFill(start_color=rojo_malo,  end_color=rojo_malo,  fill_type="solid")
    fill_zebra   = PatternFill(start_color="F4F7FB",  end_color="F4F7FB",  fill_type="solid")
    th    = Side(border_style="thin", color="BDC3C7")
    borde = Border(left=th, right=th, top=th, bottom=th)

    # ── Encabezado institucional ──────────────────────────────────────
    ws["B2"] = "DIVISIÓN DE TECNOLOGÍAS DE LA INFORMACIÓN — EDIFICIO D5"
    ws["B2"].font = f_titulo
    ws["B3"] = "Inventario de Equipos y Mobiliario de Laboratorios"
    ws["B3"].font = f_sub
    if titulo_extra:
        ws["B4"] = titulo_extra
        ws["B4"].font = f_sub

    ws["B6"] = "Exportado:"; ws["B6"].font = f_bold
    ws["C6"] = datetime.now().strftime("%Y-%m-%d %H:%M"); ws["C6"].font = f_body
    ws["E6"] = "Total registros:"; ws["E6"].font = f_bold
    ws["F6"] = len(dataframe); ws["F6"].font = f_body

    # ── Leyenda de colores ────────────────────────────────────────────
    ws["B7"] = "Leyenda:"; ws["B7"].font = f_bold
    celdas_leyenda = [("C7", "BUENO", fill_bueno), ("D7", "REGULAR", fill_regular), ("E7", "MALO", fill_malo)]
    for celda, texto, fill in celdas_leyenda:
        ws[celda] = texto
        ws[celda].font = f_bold
        ws[celda].fill = fill
        ws[celda].border = borde
        ws[celda].alignment = Alignment(horizontal="center")

    # ── Cabeceras de tabla ────────────────────────────────────────────
    headers = ["ID", "Descripción", "Marca", "Serie", "Estado",
               "Modelo", "N° Inventario", "Laboratorio", "Categoría", "Características"]

    fila_hdr = 9
    for ci, h in enumerate(headers, start=2):
        cell = ws.cell(row=fila_hdr, column=ci, value=h)
        cell.font    = f_header
        cell.fill    = fill_hdr
        cell.border  = borde
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila_hdr].height = 28

    # ── Datos ─────────────────────────────────────────────────────────
    col_names = ["ID", "Descripción", "Marca", "Serie", "Estado",
                 "Modelo", "N° Inventario", "Laboratorio", "Categoría", "Características"]

    fila_data = fila_hdr + 1
    for idx, row in dataframe.iterrows():
        estado = str(row.get("Estado", "")).upper()
        if estado == "BUENO":
            fill_fila = fill_bueno
        elif estado == "REGULAR":
            fill_fila = fill_regular
        elif estado == "MALO":
            fill_fila = fill_malo
        else:
            fill_fila = fill_zebra if idx % 2 == 0 else None

        for ci, col in enumerate(col_names, start=2):
            valor = row.get(col, "")
            # Reemplazar "—" de visualización por vacío en el Excel
            if valor == "—":
                valor = ""
            cell = ws.cell(row=fila_data, column=ci, value=valor)
            cell.font   = f_body
            cell.border = borde
            cell.alignment = Alignment(
                horizontal="center" if ci in (2, 5, 6, 7, 9) else "left",
                vertical="center",
                wrap_text=True,
            )
            if fill_fila:
                cell.fill = fill_fila
        ws.row_dimensions[fila_data].height = 18
        fila_data += 1

    # ── Firma de responsable ──────────────────────────────────────────
    fila_firma = fila_data + 2
    ws.cell(row=fila_firma,     column=2).value = "________________________________________________"
    ws.cell(row=fila_firma,     column=2).font  = f_body
    ws.cell(row=fila_firma + 1, column=2).value = "NOMBRE Y FIRMA DEL LABORATORISTA RESPONSABLE"
    ws.cell(row=fila_firma + 1, column=2).font  = f_bold

    # ── Ajustar anchos de columna ─────────────────────────────────────
    anchos_fijos = {
        2: 8,   # ID
        3: 40,  # Descripción
        4: 18,  # Marca
        5: 22,  # Serie
        6: 12,  # Estado
        7: 16,  # Modelo
        8: 22,  # N° Inventario
        9: 14,  # Laboratorio
        10: 32, # Categoría
        11: 35, # Características
    }
    for col_idx, ancho in anchos_fijos.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
