# database/migrar_inventario.py
# ─────────────────────────────────────────────────────────────────────
# Script de migración único — crea la tabla inventario en PostgreSQL
# e importa los 446 registros del archivo INVENTARIO.xlsx.
#
# USO (una sola vez desde la raíz del proyecto):
#   python -m database.migrar_inventario
#
# El script es idempotente: si la tabla ya existe, no la recrea.
# Si los registros ya fueron importados (total > 0), omite la carga.
# ─────────────────────────────────────────────────────────────────────
import os
import sys
from openpyxl import load_workbook
from database.database import get_connection


# Ruta al Excel — ajusta si está en otra ubicación
EXCEL_PATH = os.path.join(os.path.dirname(__file__), "..", "INVENTARIO.xlsx")

# Laboratorios normalizados: mapeo de los valores del Excel a los
# nombres limpios que usa la app en settings.py
LABORATORIO_MAP = {
    "LAB_D5104":        "104",
    "LAB_D5104_RACK":   "104_RACK",
    "LAB_D5204":        "204",
    "LAB_D5211":        "211",
    "LAB_D5212":        "212",
    "LAB_D5212_NOK":    "212_NOK",
    "LAB_D5212 _NOK":   "212_NOK",
    "D5202_RACK":       "202_RACK",
    "D5212_RACK":       "212_RACK",
    "LCD5204 V.22":     "204_V22",
    "LCD5204_Rack":     "204_RACK",
    "D5 ALMACEN DE LABORATORIOS": "ALMACEN",
    "N/A":              "N/A",
}


def limpiar(valor, nullable=True) -> str | None:
    """Convierte celdas vacías / S/M / S/S / N/A a None o al string limpio."""
    if valor is None:
        return None
    s = str(valor).strip()
    if s in ("", "N/A", "S/M", "S/S", "s/m", "s/s"):
        return None if nullable else s
    return s


def crear_tabla(conn):
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS inventario (
            id              SERIAL PRIMARY KEY,
            descripcion     TEXT NOT NULL,
            marca           TEXT,
            serie           TEXT,
            estado          TEXT,          -- EF: BUENO / REGULAR / MALO
            modelo          TEXT,
            num_inventario  TEXT,          -- Número de inventario institucional
            laboratorio     TEXT,          -- Normalizado: '204', '211', etc.
            categoria       TEXT,
            caracteristicas TEXT,
            -- Campos de auditoría
            activo          BOOLEAN DEFAULT TRUE,
            creado_por      TEXT,
            fecha_creacion  TEXT,
            modificado_por  TEXT,
            fecha_modificacion TEXT
        )
    """)
    conn.commit()
    print("✅ Tabla 'inventario' verificada/creada.")


def importar_excel(conn):
    c = conn.cursor()

    # No reimportar si ya hay datos
    c.execute("SELECT COUNT(*) FROM inventario")
    if c.fetchone()[0] > 0:
        print("ℹ️  La tabla ya tiene datos. Importación omitida.")
        return

    if not os.path.exists(EXCEL_PATH):
        print(f"❌ No se encontró el archivo Excel en: {EXCEL_PATH}")
        print("   Coloca INVENTARIO.xlsx en la raíz del proyecto y vuelve a correr.")
        sys.exit(1)

    wb = load_workbook(EXCEL_PATH, read_only=True)
    ws = wb.active
    insertados = 0
    omitidos = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        descripcion = limpiar(row[0], nullable=False)
        if not descripcion:
            omitidos += 1
            continue

        marca          = limpiar(row[1])
        serie          = limpiar(row[2])
        estado         = limpiar(row[3], nullable=False) or "BUENO"
        modelo         = limpiar(row[4])
        num_inventario = limpiar(row[5])
        lab_raw        = limpiar(row[6], nullable=False) or "N/A"
        laboratorio    = LABORATORIO_MAP.get(lab_raw, lab_raw)
        categoria      = limpiar(row[7])
        caracteristicas = limpiar(row[8])

        c.execute("""
            INSERT INTO inventario
                (descripcion, marca, serie, estado, modelo, num_inventario,
                 laboratorio, categoria, caracteristicas, creado_por, fecha_creacion)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW()::TEXT)
        """, (descripcion, marca, serie, estado, modelo, num_inventario,
              laboratorio, categoria, caracteristicas, "MIGRACIÓN_EXCEL"))
        insertados += 1

    conn.commit()
    wb.close()
    print(f"✅ Importación completada: {insertados} registros insertados, {omitidos} filas vacías omitidas.")


if __name__ == "__main__":
    conn = get_connection()
    crear_tabla(conn)
    importar_excel(conn)
    conn.close()
    print("\n✅ Migración finalizada. Ya puedes usar el módulo de Inventario en la app.")
